#!/usr/bin/env python3
"""
summary_data.py — 从Output出清结果提取汇总CSV
================================================
生成3个文件供后续步骤使用:
  - summary_timeseries.csv : 91天×96时段的电价/负荷/风光/水电时序数据
  - summary_daily.csv      : 91天的日度汇总(供需比/均价/中标数等)
  - summary_revenue.csv    : 91天×8电厂的收益数据

运行:
    cd ROTS44_delivery/code
    python summary_data.py

前置条件:
    - Case/ROTS 44/Output/ 目录下有91天(20260101~20260401)的出清结果
    - 每天有 总体_out.xlsx 文件
"""
import os, csv
import openpyxl
import numpy as np

# ═══════════════════════════════════════
#  配置
# ═══════════════════════════════════════
# 自动计算相对路径
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))  # code/
DELIVERY_DIR = os.path.dirname(SCRIPT_DIR)  # ROTS44_delivery/
PROJECT_ROOT = os.path.dirname(DELIVERY_DIR)  # 项目根目录
CASE_DIR = os.path.join(PROJECT_ROOT, 'Case', 'ROTS 44')  # Case/ROTS 44 (数据源)
OUTPUT_BASE = os.path.join(DELIVERY_DIR, 'ROTS 44')  # ROTS44_delivery/ROTS 44 (输出)

# 输入目录（只读）
INPUT_DIR = os.path.join(CASE_DIR, 'Output')
# 输出目录（读写）
OUTPUT_DIR = OUTPUT_BASE

SPD = 96  # 每天96个15分钟时段


def pf(msg):
    print(msg, flush=True)


def main():
    pf("=" * 70)
    pf("  提取汇总数据 (Output → CSV)")
    pf("=" * 70)
    pf(f"  输入目录: {INPUT_DIR}")
    pf(f"  输出目录: {OUTPUT_DIR}")

    # 确保输出目录存在
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    days = sorted([d for d in os.listdir(INPUT_DIR) if d.startswith('2026')])
    pf(f"  天数: {len(days)}, 范围: {days[0]} ~ {days[-1]}")

    ts_rows  = []   # 时序数据
    rev_rows = []   # 收益数据
    daily_rows = [] # 日度汇总

    for di, day in enumerate(days):
        out_file = os.path.join(INPUT_DIR, day, '总体_out.xlsx')
        try:
            wb = openpyxl.load_workbook(out_file, read_only=True, data_only=True)
        except Exception as e:
            pf(f"  [SKIP] {day}: {e}")
            continue

        # ─── 1. 出清电价 (LevelC_电能量市场出清电价) ───
        price = [0.0] * SPD
        if 'LevelC_电能量市场出清电价' in wb.sheetnames:
            ws = wb['LevelC_电能量市场出清电价']
            for row in ws.iter_rows(min_row=4, values_only=True):
                if row[0] == 'Bus_1' and row[1] and '电能量' in str(row[1]):
                    price = [float(x) if isinstance(x, (int, float)) else 0
                             for x in row[2:2+SPD]]
                    break

        # ─── 2. 电力平衡 (LevelC_电力电量平衡表单) ───
        balance = {}
        if 'LevelC_电力电量平衡表单' in wb.sheetnames:
            ws2 = wb['LevelC_电力电量平衡表单']
            for row in ws2.iter_rows(min_row=4, values_only=True):
                name = row[0]
                if name and name in ['火电功率', '风电消纳功率', '光伏消纳功率',
                                     '总负荷功率', '库容式水电功率']:
                    vals = [float(x) if isinstance(x, (int, float)) else 0
                            for x in row[1:1+SPD]]
                    balance[str(name)] = vals

        # 构建时序行
        for t in range(SPD):
            load    = balance.get('总负荷功率', [0]*SPD)[t] * 100   # 100MW→MW
            thermal = balance.get('火电功率', [0]*SPD)[t] * 100
            wind    = balance.get('风电消纳功率', [0]*SPD)[t] * 100
            solar   = balance.get('光伏消纳功率', [0]*SPD)[t] * 100
            hydro   = balance.get('库容式水电功率', [0]*SPD)[t] * 100
            supply  = thermal + wind + solar + hydro
            sdr     = supply / load if load > 0.001 else 1.0

            ts_rows.append([day, t, price[t], load, thermal, wind, solar, hydro, sdr])

        # ─── 3. 火电厂收益 (LevelB_火电厂电能量市场收益) ───
        if 'LevelB_火电厂电能量市场收益' in wb.sheetnames:
            ws3 = wb['LevelB_火电厂电能量市场收益']
            for row in ws3.iter_rows(min_row=4, values_only=True):
                if row[0] and isinstance(row[0], str) and row[0].startswith('ThermalPlant'):
                    plant_id   = str(row[0]).strip()
                    plant_name = str(row[1]).strip() if row[1] else plant_id
                    op_cost    = float(row[2]) if isinstance(row[2], (int, float)) else 0
                    start_cost = float(row[3]) if isinstance(row[3], (int, float)) else 0
                    bid_qty    = float(row[5]) if len(row) > 5 and isinstance(row[5], (int, float)) else 0
                    bid_avg    = float(row[6]) if len(row) > 6 and isinstance(row[6], (int, float)) else 0
                    bid_rev    = float(row[7]) if len(row) > 7 and isinstance(row[7], (int, float)) else 0
                    net_rev    = float(row[8]) if len(row) > 8 and isinstance(row[8], (int, float)) else 0

                    rev_rows.append([day, plant_id, plant_name,
                                     op_cost, start_cost, bid_qty,
                                     bid_avg, bid_rev, net_rev])

        # ─── 4. 日度汇总 (LevelA_电能量市场出清总览) ───
        daily_data = {'day': day}
        if 'LevelA_电能量市场出清总览' in wb.sheetnames:
            ws4 = wb['LevelA_电能量市场出清总览']
            kv = {}
            for row in ws4.iter_rows(values_only=True):
                if row[0]:
                    kv[str(row[0])] = row[1]

            daily_data['avg_price']           = kv.get('成交均价(元/MWh)', 0)
            daily_data['supply_demand_ratio']  = kv.get('供需比', 0)
            daily_data['bid_thermal_count']    = kv.get('中标火电机组数目', 0)
            daily_data['total_trade_mwh']      = kv.get('总发电量(100MWh)', 0)
            daily_data['max_price']            = kv.get('最高节点电价(元/MWh)', 0)
            daily_data['min_price']            = kv.get('最低节点电价(元/MWh)', 0)
        else:
            # 从时序数据计算
            daily_data['avg_price']           = np.mean(price)
            daily_data['supply_demand_ratio']  = 0
            daily_data['bid_thermal_count']    = 0
            daily_data['total_trade_mwh']      = 0
            daily_data['max_price']            = max(price)
            daily_data['min_price']            = min(price)

        daily_rows.append(daily_data)
        wb.close()

        if (di + 1) % 15 == 0:
            pf(f"    已处理 {di+1}/{len(days)} 天")

    pf(f"    完成, 共 {len(days)} 天")

    # ═══════════════════════════════════════
    #  写入CSV
    # ═══════════════════════════════════════

    # 1. summary_timeseries.csv
    ts_path = os.path.join(OUTPUT_DIR, 'summary_timeseries.csv')
    with open(ts_path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['day', 'period', 'price', 'load_MW', 'thermal_MW',
                     'wind_MW', 'solar_MW', 'hydro_MW', 'sdr'])
        w.writerows(ts_rows)

    # 2. summary_daily.csv
    daily_path = os.path.join(OUTPUT_DIR, 'summary_daily.csv')
    if daily_rows:
        fieldnames = list(daily_rows[0].keys())
        with open(daily_path, 'w', newline='', encoding='utf-8') as f:
            w = csv.DictWriter(f, fieldnames=fieldnames)
            w.writeheader()
            w.writerows(daily_rows)

    # 3. summary_revenue.csv
    rev_path = os.path.join(OUTPUT_DIR, 'summary_revenue.csv')
    with open(rev_path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['day', 'plant_id', 'plant_name', 'op_cost', 'start_cost',
                     'bid_qty', 'bid_avg_price', 'bid_revenue', 'net_revenue'])
        w.writerows(rev_rows)

    # ═══════════════════════════════════════
    #  打印统计
    # ═══════════════════════════════════════
    prices = [r[2] for r in ts_rows]
    net_revs = [r[8] for r in rev_rows]

    pf(f"\n  输出文件:")
    pf(f"    {ts_path} ({len(ts_rows)} rows, {os.path.getsize(ts_path)/1024:.0f} KB)")
    pf(f"    {daily_path} ({len(daily_rows)} rows)")
    pf(f"    {rev_path} ({len(rev_rows)} rows)")

    pf(f"\n  电价统计:")
    pf(f"    min={min(prices):.1f}, max={max(prices):.1f}, "
       f"mean={np.mean(prices):.1f}, median={np.median(prices):.1f}")
    pf(f"    尖峰(>500): {sum(1 for p in prices if p>500)}/{len(prices)} "
       f"({sum(1 for p in prices if p>500)/len(prices)*100:.1f}%)")

    pf(f"\n  收益统计:")
    pf(f"    总净收益: {sum(net_revs):.0f}万, 日均: {sum(net_revs)/len(days):.0f}万")
    pf(f"    亏损记录: {sum(1 for r in net_revs if r<0)}/{len(net_revs)}")

    pf(f"\n  ✓ 数据提取完成")


if __name__ == '__main__':
    main()
