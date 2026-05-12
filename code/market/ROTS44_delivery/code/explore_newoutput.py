#!/usr/bin/env python3
"""
explore_newoutput.py — 查看出清结果结构并与基线对比
运行: python explore_newoutput.py
"""
import os, sys

# 自动计算相对路径
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))  # code/
DELIVERY_DIR = os.path.dirname(SCRIPT_DIR)  # ROTS44_delivery/
PROJECT_ROOT = os.path.dirname(DELIVERY_DIR)  # 项目根目录
CASE_DIR = os.path.join(PROJECT_ROOT, 'Case', 'ROTS 44')  # Case/ROTS 44
OUTPUT_DIR = os.path.join(DELIVERY_DIR, 'ROTS 44')  # ROTS44_delivery/ROTS 44

NEW_DIR = os.path.join(OUTPUT_DIR, 'clearing_results', 'strategy_output')
OLD_DIR = os.path.join(CASE_DIR, 'Output')

# 1. 新输出文件结构
print("=" * 70)
print("新Output结构 (策略出清结果)")
print("=" * 70)
days = sorted(os.listdir(NEW_DIR))
print(f"天数: {len(days)}, 范围: {days[0]}~{days[-1]}")

for day in [days[0], days[-1]]:
    dp = os.path.join(NEW_DIR, day)
    files = os.listdir(dp)
    print(f"\n  {day}/:")
    for f in sorted(files):
        sz = os.path.getsize(os.path.join(dp, f)) / 1024
        print(f"    {f} ({sz:.0f} KB)")

# 2. 对比旧输出
print("\n" + "=" * 70)
print("旧Output结构 (基线BS3出清结果)")
print("=" * 70)
old_days = sorted([d for d in os.listdir(OLD_DIR) if d.startswith('2026')])
print(f"天数: {len(old_days)}, 范围: {old_days[0]}~{old_days[-1]}")

# 检查重叠天数
overlap = sorted(set(days) & set(old_days))
print(f"重叠天数: {len(overlap)} ({overlap[0]}~{overlap[-1]})")

# 3. 检查文件名是否一致
day0 = overlap[0]
new_files = set(os.listdir(os.path.join(NEW_DIR, day0)))
old_files = set(os.listdir(os.path.join(OLD_DIR, day0)))
print(f"\n{day0} 文件对比:")
print(f"  新: {sorted(new_files)}")
print(f"  旧: {sorted(old_files)}")
print(f"  相同: {sorted(new_files & old_files)}")
print(f"  仅新: {sorted(new_files - old_files)}")
print(f"  仅旧: {sorted(old_files - new_files)}")

# 4. 尝试读取一个xlsx看sheet结构
try:
    import openpyxl
    for f in sorted(os.listdir(os.path.join(NEW_DIR, day0))):
        if f.endswith(('.xlsx', '.xlsm')):
            fp = os.path.join(NEW_DIR, day0, f)
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
            print(f"\n  {f} sheets:")
            for sn in wb.sheetnames:
                ws = wb[sn]
                # count rows
                rc = 0
                for _ in ws.iter_rows(values_only=True):
                    rc += 1
                print(f"    {sn} ({rc} rows)")
            wb.close()
            break
except Exception as e:
    print(f"  读取失败: {e}")